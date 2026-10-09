from docx import Document
from docx.oxml.ns import qn

from engine.convert import convert
from engine.options import Options

W_TD = qn("w:textDirection")


def run(samples, tmp_path, name, **opts):
    out = tmp_path / f"{name}.docx"
    summary = convert(samples[name], out, Options.from_dict(opts or None), tmp_path / "work")
    return Document(out), summary


def body_text(doc):
    return [p.text for p in doc.paragraphs if p.text.strip()]


def section_dirs(doc):
    return [(s._sectPr.find(W_TD).get(qn("w:val")) if s._sectPr.find(W_TD) is not None else None)
            for s in doc.sections]


def test_vertical_japanese_with_furigana(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "ja_vertical_ruby")
    assert summary["directions"] == ["v"]
    assert summary["languages"] == ["ja-JP"]
    assert section_dirs(doc) == ["tbRl"]
    text = "\n".join(body_text(doc))
    assert "見当(けんとう)がつかぬ" in text
    assert "人間(にんげん)というもの" in text
    assert doc.paragraphs[0].style.name == "Heading 1" or body_text(doc)[0] == "吾輩は猫である"


def test_furigana_can_be_dropped(samples, tmp_path):
    doc, _ = run(samples, tmp_path, "ja_vertical_ruby", choices=["ja_v"], furigana="drop")
    text = "\n".join(body_text(doc))
    assert "見当がつかぬ" in text
    assert "けんとう" not in text


def test_vertical_paragraphs_run_on_across_pages(samples, tmp_path):
    doc, _ = run(samples, tmp_path, "ja_vertical_long")
    expected = [f"　第{i}段落。" + "親譲りの無鉄砲で小供の時から損ばかりしている。" * 14 for i in range(1, 9)]
    assert body_text(doc) == expected


def test_traditional_chinese_structure(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "zh_tw_horizontal")
    assert summary["languages"][0] == "zh-TW"
    styles = {p.text: p.style.name for p in doc.paragraphs if p.text}
    assert styles["臺灣的歷史與文化"] == "Heading 1"
    assert styles["第一節　概述"] == "Heading 2"
    assert styles["第一個重點項目"] == "List Bullet"
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["臺北", "250萬", "272"]
    assert len(doc.inline_shapes) == 1
    assert section_dirs(doc) == [None]


def test_two_column_reading_order(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "zh_cn_two_columns")
    assert summary["languages"] == ["zh-CN"]
    paras = body_text(doc)
    assert len(paras) == 6
    for i, p in enumerate(paras, 1):
        assert p.startswith(f"第{i}段：") and p.endswith("是否正确。")


def test_vertical_block_on_horizontal_page(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "mixed_directions")
    assert section_dirs(doc) == [None]
    cell = doc.tables[0].cell(0, 0)
    assert cell.text == "直排標語文字保持不變"
    assert cell._tc.tcPr.find(W_TD).get(qn("w:val")) == "tbRl"


def test_korean(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "ko_horizontal")
    assert summary["languages"] == ["ko-KR"]
    assert "이 문서는 한국어 텍스트가" in "\n".join(body_text(doc))


def test_forced_vertical_output(samples, tmp_path):
    doc, summary = run(samples, tmp_path, "zh_tw_horizontal", choices=["zh_tra_v"])
    assert section_dirs(doc) == ["tbRl"]
    assert summary["languages"] == ["zh-TW"]


def test_pause_resume_reuses_finished_pages(samples, tmp_path):
    work = tmp_path / "work"
    calls = []

    def stop_after_first():
        return len(calls) >= 1

    import pytest
    with pytest.raises(InterruptedError):
        convert(samples["ja_vertical_long"], tmp_path / "a.docx", Options(), work,
                progress=lambda d, t: calls.append(d), should_stop=stop_after_first)
    assert (work / "page_0001.json").exists() and not (work / "page_0002.json").exists()
    convert(samples["ja_vertical_long"], tmp_path / "a.docx", Options(), work)
    assert (tmp_path / "a.docx").exists()


def test_unicode_cmap_font_without_to_unicode(tmp_path):
    """ReportLab pairs MSung-Light (CNS1) with UniGB-UCS2-H and writes the
    text as UCS-2; without a fix every character extracts as a wrong one."""
    import pymupdf

    text = "小我：人類的現狀"
    pdf = pymupdf.open()
    page = pdf.new_page()
    font = pdf.get_new_xref()
    pdf.update_object(font, (
        "<< /Type /Font /Subtype /Type0 /BaseFont /MSung-Light /Encoding /UniGB-UCS2-H "
        "/DescendantFonts [ << /Type /Font /Subtype /CIDFontType0 /BaseFont /MSung-Light "
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (CNS1) /Supplement 1 >> /DW 1000 "
        "/FontDescriptor << /Type /FontDescriptor /FontName /MSung-Light /Flags 6 "
        "/FontBBox [ -160 -249 1015 888 ] /ItalicAngle 0 /Ascent 752 /Descent -271 "
        "/CapHeight 737 /StemV 58 >> >> ] >>"))
    pdf.xref_set_key(page.xref, "Resources", f"<< /Font << /F2 {font} 0 R >> >>")
    contents = pdf.get_new_xref()
    pdf.update_object(contents, "<<>>")
    pdf.update_stream(contents, f"BT /F2 12 Tf 72 720 Td <{text.encode('utf-16-be').hex()}> Tj ET".encode())
    pdf.xref_set_key(page.xref, "Contents", f"{contents} 0 R")
    src = tmp_path / "reportlab.pdf"
    pdf.save(src)

    out = tmp_path / "reportlab.docx"
    summary = convert(src, out, Options.from_dict(None), tmp_path / "work")
    assert text in "\n".join(body_text(Document(out)))
    assert summary["languages"] == ["zh-TW"]
