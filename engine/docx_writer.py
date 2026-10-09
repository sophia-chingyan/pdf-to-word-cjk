"""Write the analysed pages as an editable Word document.

Every PDF page starts a new Word page, with the PDF's margins, so each Word
page holds the same text as its PDF page.

Vertical pages become Word sections with text direction tbRl (top to
bottom, columns right to left). A block whose direction differs from its
page goes into a borderless one-cell table with its own text direction,
which keeps it editable without floating text boxes.
"""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .convert import CJK_LATIN_SPACE
from .layout import BULLETS
from .options import Options
from .scripts import detect_language

FONTS = {  # language -> (serif, sans), names Word knows on Windows and Mac
    "zh-TW": ("PMingLiU", "Microsoft JhengHei"),
    "zh-CN": ("SimSun", "Microsoft YaHei"),
    "ja-JP": ("Yu Mincho", "Yu Gothic"),
    "ko-KR": ("Batang", "Malgun Gothic"),
}
MARGIN = Cm(2)


def _set_east_asia(rpr, font: str, lang: str):
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), font)
    fonts.set(qn("w:hint"), "eastAsia")
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang")
        rpr.append(lang_el)
    lang_el.set(qn("w:eastAsia"), lang)


def _text_direction(section, value: str | None):
    sect = section._sectPr
    for old in sect.findall(qn("w:textDirection")):
        sect.remove(old)
    if value:
        td = OxmlElement("w:textDirection")
        td.set(qn("w:val"), value)
        anchor = None
        for tag in ("w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings"):
            anchor = sect.find(qn(tag))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(td)
        else:
            sect.append(td)


def _page_key(page: dict) -> tuple:
    return page["dir"], round(page["width"]), round(page["height"])


def _margins(pages: list[dict]) -> dict[str, float]:
    """Margins in points that fit the text area of these PDF pages.

    Left and top follow the typical page. Right and bottom take the
    smallest one seen: the widest line sets the text width, and a full PDF
    page still fits on one Word page.
    """
    boxes = []
    for p in pages:
        rs = [r["bbox"] for r in p["regions"] if r["type"] in ("text", "table")]
        if rs:
            boxes.append((min(b[0] for b in rs), min(b[1] for b in rs),
                          p["width"] - max(b[2] for b in rs), p["height"] - max(b[3] for b in rs)))
    if not boxes:
        return {side: MARGIN.pt for side in ("left", "top", "right", "bottom")}

    def median(vals):
        vals = sorted(vals)
        return vals[len(vals) // 2]

    lo, hi = Cm(0.5).pt, Cm(6).pt
    lefts, tops, rights, bottoms = zip(*boxes)
    left, top, right = median(lefts), median(tops), min(rights)
    bottom = min(min(bottoms), MARGIN.pt)
    return {side: max(lo, min(hi, v)) for side, v in
            (("left", left), ("top", top), ("right", right), ("bottom", bottom))}


def _page_setup(section, page: dict, vertical: bool, margins: dict[str, float]):
    w, h = page["width"], page["height"]
    section.orientation = WD_ORIENT.LANDSCAPE if w > h else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Pt(w), Pt(h)
    for side, value in margins.items():
        setattr(section, f"{side}_margin", Pt(value))
    _text_direction(section, "tbRl" if vertical else None)


def _one_cell(doc, direction: str, width_pt: float):
    """A borderless single-cell table whose text runs in ``direction``."""
    table = doc.add_table(rows=1, cols=1)
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tbl_pr.append(borders)
    cell = table.cell(0, 0)
    cell.width = Pt(width_pt)
    tc_pr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), direction)
    tc_pr.append(td)
    return table, cell


def _heading_levels(pages: list[dict]) -> tuple[float, dict[float, int]]:
    sizes = Counter()
    for p in pages:
        for r in p["regions"]:
            if r["type"] == "text":
                for para in r["paragraphs"]:
                    sizes[round(para["size"] * 2) / 2] += para["chars"]
    if not sizes:
        return 11.0, {}
    body = sizes.most_common(1)[0][0]
    big = sorted({s for s in sizes if s >= body * 1.15}, reverse=True)[:3]
    return body, {s: i + 1 for i, s in enumerate(big)}


def write_docx(pages: list[dict], options: Options, workdir: Path, out_path) -> dict:
    """Write the document; returns warnings and the languages and directions used."""
    allowed = options.languages
    all_text = "".join(r.get("text", "") for p in pages for r in p["regions"])
    doc_lang = detect_language(all_text, allowed)
    body, levels = _heading_levels(pages)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(body)
    # The PDF's own line pitch is set on each paragraph; no extra space
    # between paragraphs, as in typeset text.
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    _set_east_asia(normal.element.get_or_add_rPr(), FONTS[doc_lang][0], doc_lang)
    for lvl in (1, 2, 3):
        st = doc.styles[f"Heading {lvl}"]
        _set_east_asia(st.element.get_or_add_rPr(), FONTS[doc_lang][1], doc_lang)

    warnings: list[str] = []
    langs_used: Counter = Counter()
    dirs_used: Counter = Counter()
    ruby_total = 0
    current = None  # (direction, width, height) of the open section
    first = True
    # The last body paragraph written, while it may still continue in the
    # next column or page: (paragraph, size, direction).
    open_para = None
    # True until the first paragraph of a new page is written; that
    # paragraph then starts on a new Word page.
    page_break = False
    margins = {}

    def new_paragraph():
        nonlocal page_break
        p = doc.add_paragraph()
        if page_break:
            p.paragraph_format.page_break_before = True
            page_break = False
        return p

    def before_block():
        # Word ignores "page break before" inside tables, so a table that
        # opens a page gets an empty paragraph that carries the break.
        if page_break:
            new_paragraph()

    for page in pages:
        warnings.extend(page["warnings"])
        ruby_total += page.get("ruby", 0)
        key = _page_key(page)
        if key != current:
            section = doc.sections[0] if first else doc.add_section(WD_SECTION.NEW_PAGE)
            if key not in margins:
                margins[key] = _margins([p for p in pages if _page_key(p) == key])
            _page_setup(section, page, page["dir"] == "v", margins[key])
            current = key
            page_break = False
        else:
            page_break = not first
        first = False
        m = margins[key]
        text_width = page["width"] - m["left"] - m["right"]
        for region in page["regions"]:
            kind = region["type"]
            if kind != "text":
                open_para = None
            if kind == "image":
                x0, y0, x1, y1 = region["bbox"]
                path = workdir / region["file"]
                if path.exists():
                    new_paragraph().add_run().add_picture(str(path), width=Pt(min(x1 - x0, text_width)))
            elif kind == "table":
                before_block()
                rows = region["rows"]
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = "Table Grid"
                lang = detect_language("".join("".join(r) for r in rows), allowed, doc_lang)
                for i, row in enumerate(rows):
                    for j, val in enumerate(row):
                        para = table.cell(i, j).paragraphs[0]
                        run = para.add_run(CJK_LATIN_SPACE.sub("", val))
                        _set_east_asia(run._r.get_or_add_rPr(), FONTS[lang][0], lang)
            else:
                lang = detect_language(region["text"], allowed, doc_lang)
                langs_used[lang] += len(region["text"])
                dirs_used[region["dir"]] += len(region["text"])
                font = FONTS[lang][1 if region.get("sans") else 0]
                target = doc
                if region["dir"] != page["dir"]:
                    before_block()
                    x0, y0, x1, y1 = region["bbox"]
                    if region["dir"] == "v":
                        table, cell = _one_cell(doc, "tbRl", x1 - x0 + 12)
                        tr_pr = table.rows[0]._tr.get_or_add_trPr()
                        height = OxmlElement("w:trHeight")
                        height.set(qn("w:val"), str(int((y1 - y0 + 12) * 20)))
                        height.set(qn("w:hRule"), "atLeast")
                        tr_pr.append(height)
                    else:
                        table, cell = _one_cell(doc, "lrTb", y1 - y0 + 12)
                    target = cell
                if target is not doc:
                    open_para = None
                first_para = True
                for para in region["paragraphs"]:
                    level = levels.get(round(para["size"] * 2) / 2) if para["chars"] <= 80 else None
                    plain = not level and not para["list"]
                    joined = bool(target is doc and first_para and open_para and plain and not para.get("indent")
                                  and abs(open_para[1] - para["size"]) < 0.5 and open_para[2] == region["dir"])
                    if joined:
                        p = open_para[0]  # the paragraph runs on from the previous column or page
                        if page_break:
                            # Same paragraph, but its rest starts the next page.
                            p.add_run().add_break(WD_BREAK.PAGE)
                            page_break = False
                    elif target is doc:
                        p = new_paragraph()
                    elif first_para:
                        p = target.paragraphs[0]
                    else:
                        p = target.add_paragraph()
                    if not joined:
                        fmt = p.paragraph_format
                        pitch = region.get("pitch")
                        if pitch and not level and pitch > para["size"]:
                            fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                            fmt.line_spacing = Pt(pitch)
                        if plain and para.get("indent_pt"):
                            fmt.first_line_indent = Pt(para["indent_pt"])
                    first_para = False
                    runs = [list(r) for r in para["runs"]]
                    if level:
                        p.style = doc.styles[f"Heading {level}"]
                    elif para["list"] == "bullet":
                        p.style = doc.styles["List Bullet"]
                        runs[0][0] = runs[0][0].lstrip(" 　").lstrip(BULLETS).lstrip(" 　")
                    for text, bold in runs:
                        if not text:
                            continue
                        run = p.add_run(text)
                        run.bold = bold or None
                        if not level and abs(para["size"] - body) > 0.1 * body:
                            run.font.size = Pt(round(para["size"] * 2) / 2)
                        _set_east_asia(run._r.get_or_add_rPr(), font, lang)
                    open_para = (p, para["size"], region["dir"]) if (
                        target is doc and plain and para.get("open_end")) else None
                if target is not doc:
                    doc.add_paragraph()
        if page_break:
            new_paragraph()  # an empty PDF page stays an empty Word page

    doc.core_properties.title = Path(out_path).stem
    doc.save(str(out_path))
    if ruby_total:
        warnings.append(f"已處理 {ruby_total} 組注音假名（振假名）；距離較遠的注音可能未被辨識。")
    return {
        "warnings": warnings,
        "languages": [lang for lang, _ in langs_used.most_common()],
        "directions": [d for d, _ in dirs_used.most_common()],
    }
